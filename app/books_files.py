"""E-Books und Hoerbuecher aus Ordnern einlesen (Probelauf, dann Uebernahme), Titelbilder.

Erkannt werden:
  E-Books   .epub (Titel, Autor, ISBN, Reihe, Titelbild aus den Metadaten), .pdf (Info-Feld oder Dateiname),
            .mobi/.azw/.azw3 (EXTH-Kopf), .fb2, .cbz/.cbr; Calibre-Bibliotheken ueber metadata.opf und cover.jpg
  Hoerbuecher  .m4b (eine Datei = ein Buch) und Ordner mit .mp3/.m4a/.ogg/.opus/.flac (ein Ordner = ein Buch,
            Unterordner CD1/CD 2/Disc 3 gehoeren dazu). Tags per ffprobe: Album = Titel, Interpret = Autor, Komponist = Sprecher.
Danach Abgleich ueber ISBN (DNB, Open Library) oder Titel und Autor."""
import glob, json, logging, os, re, subprocess, threading, time, zipfile
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher
from . import books_db as B
from . import books_api as A
from . import books_extras as BE
from .films_art import image_ext
from . import fs
from . import users as U

EBOOK_EXT = {'.epub', '.pdf', '.mobi', '.azw', '.azw3', '.kf8', '.fb2', '.cbz', '.cbr'}
AUDIO_EXT = {'.mp3', '.m4a', '.m4b', '.ogg', '.opus', '.flac', '.aac', '.wma'}
DISC_DIR = re.compile(r'^(cd|disc|disk|teil|part)[ ._-]*\d+$', re.I)
SKIP_DIRS = {'@eadir', '#recycle', '.trash', 'lost+found', '$recycle.bin'}
COVER_NAMES = ('cover.jpg', 'cover.jpeg', 'cover.png', 'folder.jpg', 'folder.png', 'front.jpg', 'poster.jpg')
JOB = {'state': 'idle', 'folders': [], 'done': 0, 'total': 0, 'items': [], 'error': '', 'result': None, 'phase': '', 'found': 0,
       'current': '', 'audio': True, 'source': 'files'}
_lock = threading.Lock()


# ---------------------------------------------------------------- Metadaten aus Dateien
def _l(tag):
    return tag.rsplit('}', 1)[-1].lower()


def parse_opf(xml_bytes, base_dir=''):
    """OPF (EPUB oder Calibre metadata.opf) -> dict mit title, authors, isbn, language, publisher, year, series, series_no,
    description, subjects und cover_href (Pfad zum Titelbild relativ zu base_dir)."""
    root = ET.fromstring(xml_bytes)
    meta = next((e for e in root.iter() if _l(e.tag) == 'metadata'), root)
    out = {'_people': []}
    refines = {}
    for e in meta:
        if _l(e.tag) == 'meta' and e.get('refines'):
            refines.setdefault(e.get('refines').lstrip('#'), {})[e.get('property')] = (e.text or '').strip()
    isbns, subjects = [], []
    cover_id = None
    for e in meta:
        t, txt = _l(e.tag), (e.text or '').strip()
        if t == 'title' and txt and 'title' not in out:
            out['title'] = txt
        elif t == 'creator' and txt:
            role = next((v for k, v in e.attrib.items() if k.endswith('role')), None) or refines.get(e.get('id') or '', {}).get('role') or 'aut'
            out['_people'].append({'name': B_name(txt), 'role': {'aut': 'aut', 'trl': 'trl', 'ill': 'ill', 'nrt': 'nrr', 'nrr': 'nrr', 'edt': 'edt'}.get(role, 'aut')})
        elif t == 'identifier' and txt:
            i = B.norm_isbn(re.sub(r'^(urn:)?isbn:', '', txt, flags=re.I))
            if i:
                isbns.append(i)
        elif t == 'language' and txt:
            out['language'] = {'de': 'ger', 'en': 'eng', 'fr': 'fre', 'es': 'spa', 'it': 'ita', 'nl': 'dut'}.get(txt[:2].lower(), txt[:3].lower())
        elif t == 'publisher' and txt:
            out['publisher'] = txt
        elif t == 'date' and txt:
            m = re.search(r'(1[5-9]\d\d|20\d\d)', txt)
            if m and int(m.group(1)) > 1000 and not txt.startswith('0101'):
                out['year'] = int(m.group(1))
        elif t == 'description' and txt:
            out['description'] = re.sub(r'<[^>]+>', ' ', txt)
            out['description'] = re.sub(r'\s+', ' ', out['description']).strip()[:3000]
        elif t == 'subject' and txt:
            subjects.append(txt)
        elif t == 'meta':
            name, prop = e.get('name') or '', e.get('property') or ''
            if name == 'calibre:series':
                out['series'] = e.get('content')
            elif name == 'calibre:series_index':
                v = e.get('content') or ''
                out['series_no'] = v[:-2] if v.endswith('.0') else v
            elif name == 'cover':
                cover_id = e.get('content')
            elif prop == 'belongs-to-collection' and txt:
                out['series'] = txt
                gp = refines.get(e.get('id') or '', {}).get('group-position')
                if gp:
                    out['series_no'] = gp
    out['authors'] = '; '.join(p['name'] for p in out['_people'] if p['role'] == 'aut') or None
    out['isbn'] = isbns[0] if isbns else None
    if subjects:
        out['subjects'] = ', '.join(subjects[:8])
    for e in root.iter():  # Titelbild im Manifest
        if _l(e.tag) == 'item' and (e.get('id') == cover_id or 'cover-image' in (e.get('properties') or '')):
            out['cover_href'] = os.path.normpath(os.path.join(base_dir, e.get('href') or '')).replace('\\', '/')
            break
    return out


def B_name(n):
    """'Tolkien, J. R. R.' -> 'J. R. R. Tolkien'; sonst unveraendert."""
    n = re.sub(r'\s+', ' ', n).strip()
    if n.count(',') == 1 and not re.search(r'\b(und|and|&)\b', n):
        last, first = [x.strip() for x in n.split(',')]
        return f'{first} {last}'.strip()
    return n


def epub_meta(path):
    with fs.open(path, 'rb') as fh, zipfile.ZipFile(fh) as z:
        cont = ET.fromstring(z.read('META-INF/container.xml'))
        opf = next(e.get('full-path') for e in cont.iter() if _l(e.tag) == 'rootfile')
        d = parse_opf(z.read(opf), os.path.dirname(opf))
        if d.get('cover_href'):
            try:
                d['_cover'] = z.read(d['cover_href'])
            except KeyError:
                pass
    return d


def pdf_meta(path):
    """Titel und Autor aus dem Info-Feld, wenn es unkomprimiert lesbar ist."""
    d = {}
    try:
        with fs.open(path, 'rb') as fh:
            head = fh.read(65536)
            fh.seek(max(0, fs.getsize(path) - 65536))
            tail = fh.read()
    except OSError:
        return d
    for key, field in ((b'/Title', 'title'), (b'/Author', 'authors')):
        m = re.search(re.escape(key) + rb'\s*\(((?:[^()\\]|\\.){2,200})\)', tail + head)
        if m:
            raw = m.group(1)
            try:
                txt = raw.decode('utf-16') if raw.startswith(b'\xfe\xff') else raw.decode('latin-1')
            except UnicodeDecodeError:
                continue
            txt = re.sub(r'\\(.)', r'\1', txt).strip()
            if txt and not re.match(r'(microsoft|untitled|unbenannt|document|dokument)', txt, re.I):
                d[field] = txt
    if d.get('authors'):
        d['_people'] = [{'name': B_name(x), 'role': 'aut'} for x in re.split(r';|&| und | and ', d['authors']) if x.strip()]
        d['authors'] = '; '.join(p['name'] for p in d['_people'])
    return d


def mobi_meta(path):
    """Titel, Autor, ISBN aus dem EXTH-Kopf von MOBI/AZW."""
    d = {}
    try:
        with fs.open(path, 'rb') as fh:
            data = fh.read(200000)
        n = int.from_bytes(data[76:78], 'big')
        r0 = int.from_bytes(data[78:82], 'big')
        if data[r0 + 16:r0 + 20] != b'MOBI':
            return d
        hl = int.from_bytes(data[r0 + 20:r0 + 24], 'big')
        enc = 'utf-8' if int.from_bytes(data[r0 + 28:r0 + 32], 'big') == 65001 else 'cp1252'
        if int.from_bytes(data[r0 + 128:r0 + 132], 'big') & 0x40:
            e = r0 + 16 + hl
            if data[e:e + 4] == b'EXTH':
                cnt = int.from_bytes(data[e + 8:e + 12], 'big')
                p = e + 12
                for _ in range(min(cnt, 200)):
                    typ, ln = int.from_bytes(data[p:p + 4], 'big'), int.from_bytes(data[p + 4:p + 8], 'big')
                    val = data[p + 8:p + ln].decode(enc, 'replace').strip()
                    if typ == 100:
                        d.setdefault('_people', []).append({'name': B_name(val), 'role': 'aut'})
                    elif typ == 101:
                        d['publisher'] = val
                    elif typ == 104 and B.norm_isbn(val):
                        d['isbn'] = B.norm_isbn(val)
                    elif typ == 503:
                        d['title'] = val
                    elif typ == 106:
                        m = re.search(r'(1[5-9]\d\d|20\d\d)', val)
                        if m:
                            d['year'] = int(m.group(1))
                    p += ln
        if not d.get('title'):
            off = int.from_bytes(data[r0 + 84:r0 + 88], 'big')
            ln = int.from_bytes(data[r0 + 88:r0 + 92], 'big')
            if 0 < ln < 500:
                d['title'] = data[r0 + off:r0 + off + ln].decode(enc, 'replace').strip()
        if d.get('_people'):
            d['authors'] = '; '.join(p['name'] for p in d['_people'])
    except Exception:
        pass
    return d


STOP = re.compile(r'\b(der|die|das|des|dem|den|ein|eine|einer|mein|meine|und|im|in|am|vom|von der|the|of|and|a|an|my|le|la|les|du|de la|el|los|las|del|band|teil|vol)\b', re.I)


def name_like(s):
    """Sieht der Text wie ein Personenname aus? ('Frank Schätzing', 'Schätzing, Frank', 'J. R. R. Tolkien')"""
    s = s.strip()
    if ',' in s and len(s.split(',')) == 2:
        return True
    words = s.split()
    return 2 <= len(words) <= 4 and not STOP.search(s) and all(w[:1].isupper() for w in words) and not re.search(r'\d', s)


def name_meta(path):
    """'Autor - Titel.epub', 'Titel (Autor).pdf', 'Reihe 03 - Titel' ... aus dem Dateinamen."""
    stem = os.path.splitext(os.path.basename(path))[0]
    s = re.sub(r'[_]+', ' ', stem)
    s = re.sub(r'\s*\((?:z-lib|libgen|retail|epub|ebook)[^)]*\)', '', s, flags=re.I).strip()
    d = {}
    m = re.search(r'\b(97[89][\d-]{10,14})\b', s)
    if m and B.norm_isbn(m.group(1)):
        d['isbn'] = B.norm_isbn(m.group(1))
        s = s.replace(m.group(1), ' ').strip(' -')
    m = re.match(r'^(.*?)\s*\(([^()]+)\)\s*$', s)
    if m and not re.search(r'\d{4}', m.group(2)):
        d['title'], author = m.group(1), m.group(2)
    elif ' - ' in s:
        a, b = s.split(' - ', 1)
        na, nb = name_like(a), name_like(b)
        d['title'], author = (a, b) if (nb and not na) else (b, a)  # sonst 'Autor - Titel'
    else:
        d['title'], author = s, ''
    if author:
        d['_people'] = [{'name': B_name(x), 'role': 'aut'} for x in re.split(r'\s*(?:;|&| und | and )\s*', author) if x.strip()]
        d['authors'] = '; '.join(p['name'] for p in d['_people'])
    d['title'] = re.sub(r'\s+', ' ', d['title']).strip(' -.')
    return d


def audio_meta(files):
    """Tags und Laufzeit per ffprobe (erste Datei fuer Tags, alle fuer die Laufzeit)."""
    d, total = {}, 0
    for i, f in enumerate(files[:400]):
        try:
            out = subprocess.run(['ffprobe', '-v', 'quiet', '-print_format', 'json', '-show_format', fs.media_url(f)],
                                 capture_output=True, timeout=20).stdout
            fmt = json.loads(out or b'{}').get('format') or {}
        except Exception:
            continue
        total += int(float(fmt.get('duration') or 0))
        if i == 0:
            tags = {k.lower(): v for k, v in (fmt.get('tags') or {}).items()}
            d['title'] = tags.get('album') or tags.get('title')
            author = tags.get('album_artist') or tags.get('albumartist') or tags.get('artist')
            if author:
                d['_people'] = [{'name': B_name(x), 'role': 'aut'} for x in re.split(r'\s*(?:;|&|/| und | and )\s*', author) if x.strip()]
            for x in re.split(r'\s*(?:;|&|/| und )\s*', tags.get('composer') or ''):
                if x.strip() and x.strip() != author:
                    d.setdefault('_people', []).append({'name': B_name(x), 'role': 'nrr'})
            m = re.search(r'(1[5-9]\d\d|20\d\d)', tags.get('date') or tags.get('year') or '')
            if m:
                d['year'] = int(m.group(1))
            if tags.get('genre') and tags['genre'].lower() not in ('audiobook', 'hörbuch', 'hoerbuch', 'speech', 'other'):
                d['genres'] = tags['genre']
    if d.get('_people'):
        d['authors'] = '; '.join(p['name'] for p in d['_people'] if p['role'] == 'aut') or None
    d['duration_s'] = total or None
    return d


def embedded_audio_cover(path):
    try:
        r = subprocess.run(['ffmpeg', '-v', 'quiet', '-i', fs.media_url(path), '-an', '-map', '0:v:0', '-frames:v', '1', '-f', 'image2pipe', '-vcodec', 'mjpeg', '-'],
                           capture_output=True, timeout=30)
        return r.stdout if image_ext(r.stdout) else None
    except Exception:
        return None


def folder_cover(folder):
    try:
        names = fs.listdir(folder)
    except OSError:
        names = []
    for n in names:
        if n.lower() in COVER_NAMES:
            try:
                with fs.open(os.path.join(folder, n), 'rb') as fh:
                    data = fh.read(15 * 1024 * 1024)
                if image_ext(data):
                    return data
            except OSError:
                pass
    return None


# ---------------------------------------------------------------- Sammeln
def collect(folders, audio=True, progress=None):
    """Liefert Eintraege {'kind', 'path' (Datei oder Ordner), 'files', 'fmt', 'opf'} - ein Eintrag = ein Buch."""
    items = []
    for top in folders:
        for root, dirs, files in fs.walk(top):
            dirs[:] = sorted(d for d in dirs if d.lower() not in SKIP_DIRS and not d.startswith('.'))
            if progress:
                progress(root, len(items))
            low = {f.lower(): f for f in files}
            opf = os.path.join(root, low['metadata.opf']) if 'metadata.opf' in low else None
            ebooks = sorted(f for f in files if os.path.splitext(f)[1].lower() in EBOOK_EXT and not f.startswith('.'))
            if opf and ebooks:  # Calibre: ein Buch je Ordner, bevorzugt EPUB
                ebooks.sort(key=lambda f: ['.epub', '.azw3', '.mobi', '.pdf'].index(os.path.splitext(f)[1].lower())
                            if os.path.splitext(f)[1].lower() in ('.epub', '.azw3', '.mobi', '.pdf') else 9)
                items.append({'kind': 'ebook', 'path': os.path.join(root, ebooks[0]), 'fmt': os.path.splitext(ebooks[0])[1][1:].upper(),
                              'opf': opf, 'folder': root, 'others': [os.path.join(root, f) for f in ebooks[1:]]})
            else:
                for f in ebooks:
                    items.append({'kind': 'ebook', 'path': os.path.join(root, f), 'fmt': os.path.splitext(f)[1][1:].upper(), 'folder': root})
            if not audio:
                continue
            m4b = sorted(f for f in files if f.lower().endswith('.m4b'))
            for f in m4b:
                items.append({'kind': 'audiobook', 'path': os.path.join(root, f), 'files': [os.path.join(root, f)], 'fmt': 'M4B', 'folder': root})
            tracks = sorted(f for f in files if os.path.splitext(f)[1].lower() in AUDIO_EXT - {'.m4b'} and not f.startswith('.'))
            if tracks and not DISC_DIR.match(os.path.basename(root)):
                all_files = [os.path.join(root, f) for f in tracks]
                for d in sorted(dirs):  # CD1, CD2 ... gehoeren zum Hoerbuch
                    if DISC_DIR.match(d):
                        sub = os.path.join(root, d)
                        all_files += sorted(os.path.join(sub, f) for f in fs.listdir(sub) if os.path.splitext(f)[1].lower() in AUDIO_EXT)
                items.append({'kind': 'audiobook', 'path': root, 'files': all_files, 'folder': root,
                              'fmt': os.path.splitext(tracks[0])[1][1:].upper()})
            elif not tracks and any(DISC_DIR.match(d) for d in dirs):  # nur CD-Unterordner
                all_files = []
                for d in sorted(dirs):
                    if DISC_DIR.match(d):
                        sub = os.path.join(root, d)
                        all_files += sorted(os.path.join(sub, f) for f in fs.listdir(sub) if os.path.splitext(f)[1].lower() in AUDIO_EXT)
                if all_files:
                    items.append({'kind': 'audiobook', 'path': root, 'files': all_files, 'folder': root,
                                  'fmt': os.path.splitext(all_files[0])[1][1:].upper()})
    return items


def read_meta(it):
    """Metadaten eines gesammelten Eintrags (ohne Netz)."""
    d = {}
    p = it['path']
    try:
        if it.get('opf'):
            with fs.open(it['opf'], 'rb') as fh:
                d = parse_opf(fh.read())
        elif it['kind'] == 'ebook' and p.lower().endswith('.epub'):
            d = epub_meta(p)
        elif it['kind'] == 'ebook' and p.lower().endswith('.pdf'):
            d = pdf_meta(p)
        elif it['kind'] == 'ebook' and os.path.splitext(p)[1].lower() in ('.mobi', '.azw', '.azw3', '.kf8'):
            d = mobi_meta(p)
        elif it['kind'] == 'audiobook':
            d = audio_meta(it['files'])
    except Exception:
        logging.info('Metadaten nicht lesbar: %s', p)
    if it['kind'] == 'audiobook' and not d.get('title'):
        d = B_merge(d, name_meta(p if fs.isdir(p) else os.path.splitext(p)[0]))
    fn = name_meta(p)
    d = B_merge(d, fn)
    d.pop('cover_href', None)
    return d


def B_merge(a, b):
    out = dict(a or {})
    for k, v in (b or {}).items():
        if out.get(k) in (None, '', []) and v not in (None, '', []):
            out[k] = v
    return out


# ---------------------------------------------------------------- Abgleich
def norm(t):
    t = (t or '').lower().replace('ä', 'ae').replace('ö', 'oe').replace('ü', 'ue').replace('ß', 'ss')
    t = re.sub(r'\b(der|die|das|the|le|la|les|el|ein|eine|a|an)\b', ' ', t)
    return re.sub(r'[^a-z0-9]+', ' ', t).strip()


def sim(a, b):
    a, b = norm(a), norm(b)
    if not a or not b:
        return 0.0
    if a == b or (len(a) > 8 and (a in b or b in a)):
        return 1.0
    return SequenceMatcher(None, a, b).ratio()


def identify(meta):
    """Netz-Abgleich. Gibt {'status','cands','pick','note'}; pick ist das gewaehlte Ergebnis (dict) oder None."""
    if meta.get('isbn'):
        d, err = A.lookup_isbn(meta['isbn'], with_description=False)
        if d:
            return {'status': 'match', 'cands': [d], 'pick': 0, 'note': 'ISBN aus der Datei'}
    if not meta.get('title'):
        return {'status': 'nf', 'cands': [], 'pick': None, 'note': 'Kein Titel erkennbar'}
    author = (meta.get('authors') or '').split(';')[0].strip()
    last = author.split()[-1] if author else ''
    cands, err = A.search(title=meta['title'], author=last)
    if not cands and last:
        cands, err = A.search(title=meta['title'])
    if err and not cands:
        return {'status': 'nf', 'cands': [], 'pick': None, 'note': err}
    if not cands:
        return {'status': 'nf', 'cands': [], 'pick': None, 'note': 'Nicht gefunden'}
    def score(c):
        s = sim(meta['title'], c.get('title'))
        if author:
            s = s * 0.75 + 0.25 * max([sim(author, a) for a in (c.get('authors') or '').split(';')] + [0])
        if meta.get('language') and c.get('language') == meta['language']:
            s += 0.03
        return s
    ranked = sorted(range(len(cands[:10])), key=lambda i: -score(cands[i]))
    best = ranked[0]
    sc = score(cands[best])
    cands = [cands[i] for i in ranked]
    return {'status': 'match' if sc >= 0.85 else 'unsure', 'cands': cands[:6], 'pick': 0, 'score': round(sc, 2),
            'note': '' if sc >= 0.85 else 'Titel nur ähnlich'}


# ---------------------------------------------------------------- Hintergrundlauf
def scan_worker(folders, audio=True):
    try:
        with B.db() as c:
            known = B.known_paths(c)
        JOB.update(phase='collect', found=0, current='', total=0, done=0, source='files')
        with fs.unreadable() as bad:
            items = collect(folders, audio, lambda root, n: JOB.update(current=root, found=n))
        JOB['unreadable'] = bad
        items = [it for it in items if it['path'] not in known]
        JOB.update(phase='read', total=len(items), done=0, found=len(items))
        for it in items:
            JOB['current'] = os.path.basename(it['path'])
            it['meta'] = read_meta(it)
            it['meta'].pop('_cover', None)  # Bild erst bei der Uebernahme lesen
            JOB['done'] += 1
        JOB.update(phase='match', done=0)
        out = []
        with B.db() as c:
            have = {r['isbn']: (r['id'], r['title']) for r in c.execute('SELECT id, isbn, title FROM books WHERE isbn IS NOT NULL')}
        for i, it in enumerate(items):
            JOB['current'] = it['meta'].get('title') or os.path.basename(it['path'])
            try:
                r = identify(it['meta'])
            except Exception as e:
                logging.exception('Buch nicht erkannt: %s', it['path'])
                r = {'status': 'nf', 'cands': [], 'pick': None, 'note': str(e)}
            for cd in r['cands']:
                cd.pop('_isbns', None)
            r.update(idx=i, kind=it['kind'], path=it['path'], fmt=it['fmt'], meta=it['meta'], files=it.get('files'),
                     opf=it.get('opf'), others=it.get('others'))
            pick = r['cands'][0] if r['cands'] else None
            if pick and pick.get('isbn') in have and r['status'] == 'match':
                r['status'] = 'known'
                r['book_id'], r['book_title'] = have[pick['isbn']]
            if r['status'] in ('match', 'known'):
                r['default'] = 'c:0'
            elif r['status'] == 'unsure':
                r['default'] = 'c:0' if r.get('score', 0) >= 0.7 else 'plain'
            else:
                r['default'] = 'plain'  # eigene Datei: lieber mit Dateiangaben anlegen als ueberspringen
            out.append(r)
            JOB['done'] += 1
        JOB.update(items=out, state='ready', current='')
    except Exception as e:
        logging.exception('Buch-Scan fehlgeschlagen')
        JOB.update(state='error', error=str(e))


def file_size(it):
    try:
        if it['kind'] == 'audiobook':
            return sum(fs.getsize(f) for f in it.get('files') or [])
        return fs.getsize(it['path'])
    except OSError:
        return None


def save_cover(c, images_dir, bid, data):
    ext = image_ext(data)
    if not ext:
        return None
    for f in glob.glob(os.path.join(images_dir, f'b{bid}-cover*')):
        try:
            os.remove(f)
        except OSError:
            pass
    name = f'b{bid}-cover-{int(time.time() * 1000) % 10**10}.{ext}'
    with open(os.path.join(images_dir, name), 'wb') as fh:
        fh.write(data)
    B.update(c, bid, {'cover': name})
    return name


def remove_covers(images_dir, bid):
    for f in glob.glob(os.path.join(images_dir, f'b{bid}-*')):
        try:
            os.remove(f)
        except OSError:
            pass


def local_cover(it):
    """Titelbild aus der Datei oder dem Ordner (ohne Netz)."""
    p = it['path']
    try:
        if it.get('opf'):
            data = folder_cover(os.path.dirname(it['opf']))
            if data:
                return data
        if it['kind'] == 'ebook' and p.lower().endswith('.epub'):
            return epub_meta(p).get('_cover')
        if it['kind'] == 'audiobook':
            return folder_cover(p if fs.isdir(p) else os.path.dirname(p)) or embedded_audio_cover((it.get('files') or [p])[0])
        return folder_cover(os.path.dirname(p)) if it.get('opf') else None
    except Exception:
        return None


def save_book(c, data, images_dir, cover_bytes=None, kind_hint=None):
    """Buch aus Netzdaten anlegen (oder vorhandenes per ISBN zurueckgeben). Gibt (book_id, neu)."""
    if data.get('isbn'):
        old = B.by_isbn(c, data['isbn'])
        if old:
            return old[0]['id'], False
    bid = B.insert(c, data)
    img = cover_bytes or A.fetch_cover(data.get('_cover_id'), data.get('isbn') if not data.get('_cover_id') else None)
    if img:
        save_cover(c, images_dir, bid, img)
    return bid, True


def apply_worker(choices, images_dir):
    res = {'new': 0, 'attached': 0, 'plain': 0, 'skipped': 0}
    try:
        JOB.update(state='applying', total=len(JOB['items']), done=0, phase='apply')
        for it in JOB['items']:
            ch = choices.get(it['idx'], it['default'])
            JOB['current'] = it['meta'].get('title') or os.path.basename(it['path'])
            try:
                if ch == 'skip':
                    res['skipped'] += 1
                    continue
                cover = local_cover(it)
                meta = dict(it['meta'])
                if ch.startswith('c:') and ch[2:].isdigit() and int(ch[2:]) < len(it['cands']):
                    d = dict(it['cands'][int(ch[2:])])
                    if d.get('isbn') and d.get('source') == 'openlibrary' and not d.get('dnb_id') and not d.get('publisher'):
                        full, _ = A.lookup_isbn(d['isbn'], with_description=True)  # vollstaendige Angaben zur gewaehlten Ausgabe
                        d = A.merge(full, d) if full else d
                    for k in ('series', 'series_no', 'duration_s'):  # aus der Datei, wenn das Netz nichts hat
                        if meta.get(k) and not d.get(k):
                            d[k] = meta[k]
                    if it['kind'] == 'audiobook' and meta.get('_people'):  # Sprecher aus den Tags dazu
                        d['_people'] = list(d.get('_people') or []) + [p for p in meta['_people'] if p['role'] == 'nrr']
                    with B.db() as c:
                        bid, new = save_book(c, d, images_dir, cover)
                        res['new' if new else 'attached'] += 1
                else:
                    meta['source'] = 'file'
                    with B.db() as c:
                        bid = B.insert(c, meta)
                        if cover:
                            save_cover(c, images_dir, bid, cover)
                    res['plain'] += 1
                with B.db() as c:
                    if not c.execute('SELECT 1 FROM copies WHERE path=?', (it['path'],)).fetchone():
                        B.add_copy(c, bid, {'kind': it['kind'], 'path': it['path'], 'file_format': it['fmt'], 'size_bytes': file_size(it),
                                            'duration_s': meta.get('duration_s'), 'isbn': meta.get('isbn')})
                    if meta.get('duration_s'):
                        B.update(c, bid, {'duration_s': meta['duration_s']}, only_empty=True)
            except Exception:
                logging.exception('Buch nicht übernommen: %s', it.get('path'))
                res['skipped'] += 1
            finally:
                JOB['done'] += 1
        JOB.update(state='done', result=res, items=[], current='')
        BE.start()  # Zusammenfassungen und Bewertungen im Hintergrund nachladen
    except Exception as e:
        logging.exception('Buch-Import fehlgeschlagen')
        JOB.update(state='error', error=str(e))


def start(target, *args):
    with _lock:
        if JOB['state'] in ('scanning', 'applying'):
            return False
        JOB['state'] = 'scanning' if target is scan_worker else 'applying'
    U.start_thread(target, *args)  # mit dem angemeldeten Nutzer (Besitzer neuer Exemplare)
    return True


# ---------------------------------------------------------------- Liste oder Tabelle (ISBN-Liste, CSV, Excel)
HEAD = {'isbn': ('isbn', 'ean', 'barcode', 'isbn13', 'isbn-13', 'isbn10'), 'title': ('titel', 'title', 'buchtitel', 'name'),
        'authors': ('autor', 'autoren', 'author', 'authors', 'verfasser', 'auteur', 'autor/in'),
        'kind': ('format', 'art', 'typ', 'einband', 'binding', 'type'), 'status': ('gelesen', 'read', 'status', 'lu', 'leído'),
        'rating': ('bewertung', 'rating', 'note', 'sterne', 'stars'), 'notes': ('notiz', 'notizen', 'bemerkung', 'notes', 'kommentar'),
        'series': ('reihe', 'serie', 'series'), 'series_no': ('band', 'nr', 'nummer', 'volume', 'tome'), 'year': ('jahr', 'year', 'année')}
KIND_WORDS = {'hardcover': ('hardcover', 'gebunden', 'hc', 'hardback'), 'paperback': ('taschenbuch', 'paperback', 'tb', 'softcover', 'broschur'),
              'ebook': ('ebook', 'e-book', 'epub', 'kindle', 'pdf'), 'audiobook': ('hörbuch', 'hoerbuch', 'audiobook', 'audio', 'cd')}


def read_table(name, data):
    """Zeilen aus CSV, Excel oder einfacher Liste -> [{'isbn','title','authors',...}]."""
    rows = []
    low = (name or '').lower()
    if low.endswith(('.xlsx', '.xlsm')):
        from io import BytesIO
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(data), read_only=True, data_only=True).active
        rows = [[('' if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
    else:
        import csv, io
        text = data.decode('utf-8-sig', 'replace') if isinstance(data, bytes) else data
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            return []
        delim = max([';', ',', '\t'], key=lambda d: lines[0].count(d)) if any(d in lines[0] for d in ';,\t') else None
        rows = [[x.strip() for x in r] for r in csv.reader(io.StringIO('\n'.join(lines)), delimiter=delim)] if delim else [[ln.strip()] for ln in lines]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    cols = {}
    for i, h in enumerate(rows[0]):
        h2 = h.lower().strip()
        for key, names in HEAD.items():
            if h2 in names and key not in cols:
                cols[key] = i
    if cols:
        body = rows[1:]
    else:  # ohne Kopfzeile: Spalte mit ISBN erkennen, sonst Titel (und Autor in Spalte 2)
        body = rows
        first = rows[0]
        idx = next((i for i, v in enumerate(first) if B.norm_isbn(v)), None)
        cols = {'isbn': idx} if idx is not None else {'title': 0, 'authors': 1} if len(first) > 1 else {'title': 0}
    out = []
    for r in body:
        get = lambda k: (r[cols[k]] if k in cols and cols[k] < len(r) else '').strip()
        d = {'isbn': B.norm_isbn(get('isbn')) or None, 'title': get('title'), 'authors': get('authors') or None, 'notes': get('notes') or None,
             'series': get('series') or None, 'series_no': get('series_no') or None}
        if get('isbn') and not d['isbn'] and not d['title']:
            d['title'] = get('isbn')  # keine gueltige ISBN: als Titel versuchen
        k = get('kind').lower()
        d['kind'] = next((kk for kk, words in KIND_WORDS.items() if any(w in k for w in words)), None)
        s = get('status').lower()
        d['status'] = 'read' if s in ('1', 'x', 'ja', 'yes', 'true', 'gelesen', 'read', 'oui', 'sí', 'si') or re.match(r'\d{4}', s) else \
            'reading' if s in ('lese', 'reading', 'gerade') else ''
        d['read_at'] = B.parse_date(s) if re.match(r'\d', s) else None
        rt = get('rating')
        d['rating'] = int(float(rt.replace(',', '.'))) if re.fullmatch(r'\d+([.,]\d)?', rt or '') and 1 <= float(rt.replace(',', '.')) <= 10 else None
        y = get('year')
        d['year'] = int(y) if re.fullmatch(r'\d{4}', y or '') else None
        if d['authors']:
            d['_people'] = [{'name': B_name(x), 'role': 'aut'} for x in re.split(r'\s*(?:;|&| und | and )\s*', d['authors']) if x.strip()]
            d['authors'] = '; '.join(p['name'] for p in d['_people'])
        if d['isbn'] or d['title']:
            out.append(d)
    return out


def table_scan_worker(rows, default_kind):
    try:
        with B.db() as c:
            have = {r['isbn']: (r['id'], r['title']) for r in c.execute('SELECT id, isbn, title FROM books WHERE isbn IS NOT NULL')}
        JOB.update(phase='match', total=len(rows), done=0, found=len(rows), source='table')
        out, seen = [], set()
        for i, meta in enumerate(rows):
            JOB['current'] = meta.get('title') or meta.get('isbn') or ''
            key = meta.get('isbn') or (meta.get('title') or '') + '|' + (meta.get('authors') or '')
            if key in seen:  # doppelt in der Liste
                JOB['done'] += 1
                continue
            seen.add(key)
            try:
                r = identify(meta)
            except Exception as e:
                r = {'status': 'nf', 'cands': [], 'pick': None, 'note': str(e)}
            for cd in r['cands']:
                cd.pop('_isbns', None)
            kind = meta.get('kind') or (r['cands'][0].get('_kind') if r['cands'] else None) or default_kind
            r.update(idx=i, kind=kind, path=None, fmt='', meta=meta, files=None, opf=None, others=None)
            pick = r['cands'][0] if r['cands'] else None
            if pick and pick.get('isbn') in have and r['status'] == 'match':
                r['status'] = 'known'
                r['book_id'], r['book_title'] = have[pick['isbn']]
            r['default'] = 'c:0' if r['status'] in ('match', 'known') or (r['status'] == 'unsure' and r.get('score', 0) >= 0.7) else \
                ('plain' if meta.get('title') else 'skip')
            if r['status'] == 'known':  # dieselbe ISBN schon als Exemplar da: nicht doppelt anlegen
                with B.db() as c:
                    dup = c.execute('SELECT 1 FROM copies WHERE book_id=? AND isbn=?', (r['book_id'], pick.get('isbn'))).fetchone()
                if dup:
                    r['default'], r['note'] = 'skip', 'Schon als Exemplar vorhanden'
            out.append(r)
            JOB['done'] += 1
        JOB.update(items=out, state='ready', current='')
    except Exception as e:
        logging.exception('Buchliste nicht gelesen')
        JOB.update(state='error', error=str(e))


def table_apply_worker(choices, kinds, images_dir):
    res = {'new': 0, 'attached': 0, 'plain': 0, 'skipped': 0}
    try:
        JOB.update(state='applying', total=len(JOB['items']), done=0, phase='apply')
        for it in JOB['items']:
            ch = choices.get(it['idx'], it['default'])
            meta = it['meta']
            JOB['current'] = meta.get('title') or meta.get('isbn') or ''
            try:
                if ch == 'skip':
                    res['skipped'] += 1
                    continue
                kind = kinds.get(it['idx']) if kinds.get(it['idx']) in B.KINDS else it['kind']
                with B.db() as c:
                    if ch.startswith('c:') and ch[2:].isdigit() and int(ch[2:]) < len(it['cands']):
                        d = dict(it['cands'][int(ch[2:])])
                        for k in ('series', 'series_no'):
                            if meta.get(k) and not d.get(k):
                                d[k] = meta[k]
                        bid, new = save_book(c, d, images_dir)
                        res['new' if new else 'attached'] += 1
                    else:
                        bid = B.insert(c, dict(meta, source='manual'))
                        res['plain'] += 1
                    B.add_copy(c, bid, {'kind': kind, 'isbn': meta.get('isbn') or (it['cands'][0].get('isbn') if it['cands'] and ch.startswith('c:') else None)})
                    if meta.get('status') == 'read':
                        B.set_read_date(c, bid, meta.get('read_at'))
                    elif meta.get('status') == 'reading':
                        B.set_status(c, bid, 'reading')
                    if meta.get('rating'):
                        B.set_rating(c, bid, meta['rating'])
                    if meta.get('notes'):
                        B.update(c, bid, {'notes': meta['notes']}, only_empty=True)
            except Exception:
                logging.exception('Buch aus Liste nicht übernommen')
                res['skipped'] += 1
            finally:
                JOB['done'] += 1
        JOB.update(state='done', result=res, items=[], current='')
        BE.start()  # Zusammenfassungen und Bewertungen im Hintergrund nachladen
    except Exception as e:
        logging.exception('Buchliste nicht übernommen')
        JOB.update(state='error', error=str(e))


def start_table(target, *args):
    with _lock:
        if JOB['state'] in ('scanning', 'applying'):
            return False
        JOB['state'] = 'scanning' if target is table_scan_worker else 'applying'
    U.start_thread(target, *args)  # mit dem angemeldeten Nutzer (Besitzer neuer Exemplare)
    return True
