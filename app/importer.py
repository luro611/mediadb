"""Datei-Import: Tabellen lesen (xlsx, ods, csv) und Spalten erkennen. Keine Datenbank, kein Netz."""
import csv
import io
import re
import zipfile
import xml.etree.ElementTree as ET

MAX_ROWS = 20000

# Felder, die der Import braucht, und woran man die Spalte erkennt (kleingeschrieben, ohne Sonderzeichen)
SYN = {
    'artist': ['artist', 'interpret', 'kunstler', 'kuenstler', 'kunstlername', 'band', 'performer', 'musiker', 'albumartist', 'albuminterpret'],
    'title': ['title', 'titel', 'album', 'albumtitel', 'albumname', 'cd', 'cdtitel', 'name', 'release'],
    'barcode': ['barcode', 'ean', 'upc', 'ean13', 'gtin', 'strichcode'],
    'year': ['year', 'jahr', 'erscheinungsjahr', 'erschienen', 'date', 'datum', 'releasedate'],
}
LABELS = {'artist': 'Interpret', 'title': 'Albumtitel', 'barcode': 'Barcode (optional)', 'year': 'Jahr (optional)'}


def _key(h):
    return re.sub(r'[^a-z0-9]', '', str(h or '').lower().replace('ü', 'u').replace('ö', 'o').replace('ä', 'a').replace('ß', 'ss'))


def cell_text(v):
    """Zellwert als Text. Ganze Zahlen ohne '.0' (Barcodes!)."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_xlsx(data):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append([cell_text(c) for c in r])
        if len(rows) >= MAX_ROWS:
            break
    return rows


def _read_ods(data):
    ns = {'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0', 'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0', 'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}
    T = '{%s}' % ns['table']
    O = '{%s}' % ns['office']
    X = '{%s}' % ns['text']
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read('content.xml'))
    sheet = root.find('.//table:table', ns)
    rows = []
    for tr in sheet.iter(T + 'table-row'):
        rep = int(tr.get(T + 'number-rows-repeated', '1'))
        cells = []
        for tc in tr:
            if tc.tag not in (T + 'table-cell', T + 'covered-table-cell'):
                continue
            crep = min(int(tc.get(T + 'number-columns-repeated', '1')), 200)
            vt = tc.get(O + 'value-type')
            if vt in ('float', 'percentage', 'currency') and tc.get(O + 'value') is not None:
                val = cell_text(float(tc.get(O + 'value')))
            else:
                val = '\n'.join(''.join(p.itertext()) for p in tc.findall(X + 'p')).strip()
            cells.extend([val] * crep)
        while cells and cells[-1] == '':
            cells.pop()
        if not cells:
            continue            # leere Zeilen (auch lange Wiederholungen) ueberspringen
        for _ in range(min(rep, 5)):
            rows.append(cells)
        if len(rows) >= MAX_ROWS:
            break
    return rows


def _read_csv(data):
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=';,\t|').delimiter
    except csv.Error:
        delim = ';' if sample.count(';') >= sample.count(',') else ','
    rows = []
    for r in csv.reader(io.StringIO(text), delimiter=delim):
        rows.append([c.strip() for c in r])
        if len(rows) >= MAX_ROWS:
            break
    return rows


def read_table(filename, data):
    """Gibt (zeilen, fehler) zurueck. Zeilen = Liste von Listen mit Text."""
    name = (filename or '').lower()
    try:
        if name.endswith(('.xlsx', '.xlsm')):
            rows = _read_xlsx(data)
        elif name.endswith('.ods'):
            rows = _read_ods(data)
        elif name.endswith(('.csv', '.txt', '.tsv')):
            rows = _read_csv(data)
        else:
            return [], 'Dateityp nicht unterstuetzt. Erlaubt sind .xlsx, .ods und .csv.'
    except Exception as e:
        return [], f'Die Datei konnte nicht gelesen werden ({e.__class__.__name__}).'
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], 'Die Datei enthaelt keine Zeilen.'
    return rows, ''


def col_letter(i):
    s = ''
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def guess_mapping(headers):
    """Spalte je Feld anhand der Ueberschrift erraten. Gibt {feld: spaltenindex oder None} zurueck."""
    keys = [_key(h) for h in headers]
    out = {}
    used = set()
    for field, names in SYN.items():
        pick = None
        for n in names:                         # genaue Treffer nach Prioritaet der Namen
            for i, k in enumerate(keys):
                if i not in used and k == n:
                    pick = i
                    break
            if pick is not None:
                break
        out[field] = pick
        if pick is not None:
            used.add(pick)
    return out


BACKUP_FIELDS = {'barcode', 'title', 'artist', 'release_mbid', 'release_group_mbid', 'ownership', 'file_formats', 'source', 'created_at', 'tracks', 'genre', 'label', 'catalog_number', 'cover_path'}


def is_backup(headers):
    """Erkennt die eigene Exportdatei (viele Feldnamen der Datenbank als Ueberschriften)."""
    return len({str(h).strip().lower() for h in headers} & BACKUP_FIELDS) >= 8
